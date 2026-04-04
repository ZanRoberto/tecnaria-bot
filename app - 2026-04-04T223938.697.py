"""
ORACOLO COVOLO - FLASK VERSION
===============================
Senza FastAPI/Uvicorn - Solo Flask
Niente problemi di dipendenze!
"""

import os
import json
import sqlite3
from typing import List, Dict, Any
from datetime import datetime
from enum import Enum

from flask import Flask, render_template_string, request, jsonify
import httpx

from macrorule_engine import MacroruleEngine
from narrator_system import NarratorSystem

# ============================================================
# CONFIG
# ============================================================

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DATA_DIR = os.path.join(BASE_DIR, "data")
UPLOADS_DIR = os.path.join(DATA_DIR, "uploads")
DB_PATH = os.path.join(DATA_DIR, "oracolo_covolo.db")
MACRORULE_FILE = os.path.join(BASE_DIR, "macroregole_covolo_universe.json")

DEEPSEEK_API_KEY = os.getenv("DEEPSEEK_API_KEY", "").strip()
DEEPSEEK_API_URL = "https://api.deepseek.com/chat/completions"
DEEPSEEK_MODEL = "deepseek-chat"

os.makedirs(UPLOADS_DIR, exist_ok=True)
os.makedirs(DATA_DIR, exist_ok=True)

# ============================================================
# INITIALIZE SYSTEMS
# ============================================================

macrorule_engine = MacroruleEngine(MACRORULE_FILE)
narrator = NarratorSystem()

# ============================================================
# DATABASE
# ============================================================

def init_db():
    conn = sqlite3.connect(DB_PATH)
    c = conn.cursor()
    
    c.execute('''CREATE TABLE IF NOT EXISTS documents
                 (id INTEGER PRIMARY KEY,
                  filename TEXT UNIQUE,
                  upload_date TIMESTAMP,
                  file_size INTEGER,
                  content TEXT,
                  preview TEXT)''')
    
    c.execute('''CREATE TABLE IF NOT EXISTS queries
                 (id INTEGER PRIMARY KEY,
                  question TEXT,
                  answer TEXT,
                  documents_used TEXT,
                  timestamp TIMESTAMP)''')
    
    conn.commit()
    conn.close()

init_db()

# ============================================================
# FLASK APP
# ============================================================

app = Flask(__name__)
app.config['MAX_CONTENT_LENGTH'] = 50 * 1024 * 1024  # 50MB max

# ============================================================
# FILE PROCESSING
# ============================================================

def extract_text_from_file(file_path: str) -> str:
    ext = os.path.splitext(file_path)[1].lower()
    
    try:
        if ext == '.pdf':
            import PyPDF2
            text = ""
            with open(file_path, 'rb') as f:
                reader = PyPDF2.PdfReader(f)
                for page_num, page in enumerate(reader.pages, 1):
                    text += f"[PAGE {page_num}]\n"
                    text += page.extract_text() + "\n"
            return text
        
        elif ext in ['.xlsx', '.xls']:
            import openpyxl
            text = ""
            wb = openpyxl.load_workbook(file_path)
            for sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                text += f"[SHEET: {sheet_name}]\n"
                for row in ws.iter_rows(values_only=True):
                    text += " | ".join(str(cell) if cell else "" for cell in row) + "\n"
            return text
        
        elif ext in ['.docx', '.doc']:
            from docx import Document
            doc = Document(file_path)
            text = ""
            for para in doc.paragraphs:
                text += para.text + "\n"
            return text
        
        elif ext == '.txt':
            with open(file_path, 'r', encoding='utf-8') as f:
                return f.read()
        
        else:
            return "File type not supported"
    except Exception as e:
        return f"Error: {str(e)}"

# ============================================================
# DATABASE HELPERS
# ============================================================

def save_document_to_db(filename: str, content: str, file_size: int) -> bool:
    try:
        conn = sqlite3.connect(DB_PATH)
        c = conn.cursor()
        preview = content[:500] + "..." if len(content) > 500 else content
        c.execute('''INSERT OR REPLACE INTO documents 
                     (filename, upload_date, file_size, content, preview)
                     VALUES (?, ?, ?, ?, ?)''',
                  (filename, datetime.now().isoformat(), file_size, content, preview))
        conn.commit()
        conn.close()
        return True
    except Exception as e:
        print(f"Error saving: {e}")
        return False

def get_all_documents() -> List[Dict]:
    try:
        conn = sqlite3.connect(DB_PATH)
        c = conn.cursor()
        c.execute('SELECT filename, upload_date, file_size, preview FROM documents')
        docs = c.fetchall()
        conn.close()
        return [{"filename": doc[0], "upload_date": doc[1], "file_size": doc[2], "preview": doc[3][:100]}
                for doc in docs]
    except Exception:
        return []

def search_documents(query: str) -> List[Dict]:
    try:
        conn = sqlite3.connect(DB_PATH)
        c = conn.cursor()
        search_term = f"%{query.lower()}%"
        c.execute('SELECT filename, content FROM documents WHERE LOWER(content) LIKE ? OR LOWER(filename) LIKE ?',
                  (search_term, search_term))
        results = c.fetchall()
        conn.close()
        return [{"filename": result[0], "content": result[1]} for result in results]
    except Exception:
        return []

# ============================================================
# ROUTES
# ============================================================

@app.route('/')
def index():
    html = """
    <!DOCTYPE html>
    <html lang="it">
    <head>
        <meta charset="UTF-8">
        <meta name="viewport" content="width=device-width, initial-scale=1.0">
        <title>Oracolo Covolo</title>
        <style>
            * { margin: 0; padding: 0; box-sizing: border-box; }
            body { 
                font-family: -apple-system, BlinkMacSystemFont, 'Segoe UI', Roboto, sans-serif;
                background: linear-gradient(135deg, #0f172e 0%, #1a1f3a 100%);
                color: #e0e0e0;
                min-height: 100vh;
                display: flex;
            }
            .container { display: flex; width: 100%; height: 100vh; }
            .sidebar {
                width: 280px;
                background: rgba(15, 23, 46, 0.8);
                border-right: 1px solid rgba(59, 130, 245, 0.2);
                padding: 20px;
                overflow-y: auto;
            }
            .main {
                flex: 1;
                display: flex;
                flex-direction: column;
            }
            .header {
                background: rgba(59, 130, 245, 0.1);
                border-bottom: 1px solid rgba(59, 130, 245, 0.2);
                padding: 20px;
                text-align: center;
            }
            .header h1 { color: #3b82f6; font-size: 28px; margin-bottom: 5px; }
            .header p { color: #9ca3af; font-size: 14px; }
            .chat-area {
                flex: 1;
                display: flex;
                flex-direction: column;
                padding: 20px;
                overflow-y: auto;
            }
            .messages { flex: 1; overflow-y: auto; margin-bottom: 20px; }
            .message {
                margin-bottom: 15px;
                padding: 12px 15px;
                border-radius: 8px;
                max-width: 80%;
                word-wrap: break-word;
            }
            .bot-message {
                background: rgba(59, 130, 245, 0.2);
                border-left: 3px solid #3b82f6;
            }
            .user-message {
                background: rgba(168, 85, 247, 0.2);
                border-left: 3px solid #a855f7;
                margin-left: auto;
            }
            .input-area {
                display: flex;
                gap: 10px;
                margin-top: 20px;
            }
            input {
                flex: 1;
                background: rgba(30, 41, 59, 0.8);
                border: 1px solid rgba(59, 130, 245, 0.3);
                color: #e0e0e0;
                padding: 10px 15px;
                border-radius: 6px;
            }
            button {
                background: #3b82f6;
                color: white;
                border: none;
                padding: 10px 20px;
                border-radius: 6px;
                cursor: pointer;
            }
            button:hover { background: #2563eb; }
            .upload-area {
                border: 2px dashed rgba(59, 130, 245, 0.3);
                border-radius: 8px;
                padding: 15px;
                text-align: center;
                margin-bottom: 20px;
                cursor: pointer;
            }
            .file-item {
                background: rgba(59, 130, 245, 0.1);
                padding: 8px 12px;
                margin-bottom: 8px;
                border-radius: 4px;
                font-size: 12px;
                display: flex;
                justify-content: space-between;
                align-items: center;
            }
            .delete-btn {
                background: #ef4444;
                padding: 2px 8px;
                font-size: 11px;
            }
        </style>
    </head>
    <body>
        <div class="container">
            <div class="sidebar">
                <h3 style="color: #3b82f6; margin-bottom: 15px;">📁 Documenti</h3>
                <div class="upload-area" onclick="document.getElementById('file-input').click()">
                    <div>📤 Trascina o clicca</div>
                    <input type="file" id="file-input" hidden onchange="uploadFile(this)">
                </div>
                <div id="files-list"></div>
            </div>
            <div class="main">
                <div class="header">
                    <h1>🔮 Oracolo Covolo</h1>
                    <p>Sistema Intelligente - Flask Edition</p>
                </div>
                <div class="chat-area">
                    <div class="messages" id="messages"></div>
                    <div class="input-area">
                        <input type="text" id="question" placeholder="Fai una domanda..." onkeypress="if(event.key==='Enter') sendQuestion()">
                        <button onclick="sendQuestion()">Invia</button>
                    </div>
                </div>
            </div>
        </div>
        <script>
            async function sendQuestion() {
                const input = document.getElementById('question');
                const question = input.value.trim();
                if (!question) return;
                
                const messagesDiv = document.getElementById('messages');
                messagesDiv.innerHTML += `<div class="message user-message">${question}</div>`;
                input.value = '';
                messagesDiv.scrollTop = messagesDiv.scrollHeight;
                
                try {
                    const response = await fetch('/api/ask', {
                        method: 'POST',
                        headers: { 'Content-Type': 'application/json' },
                        body: JSON.stringify({ question: question })
                    });
                    const data = await response.json();
                    messagesDiv.innerHTML += `<div class="message bot-message">${data.answer || 'Errore'}</div>`;
                    messagesDiv.scrollTop = messagesDiv.scrollHeight;
                } catch (error) {
                    messagesDiv.innerHTML += `<div class="message bot-message">Errore: ${error}</div>`;
                }
            }
            
            async function uploadFile(input) {
                const file = input.files[0];
                if (!file) return;
                
                const formData = new FormData();
                formData.append('file', file);
                
                try {
                    const response = await fetch('/api/upload', {
                        method: 'POST',
                        body: formData
                    });
                    const data = await response.json();
                    if (data.status === 'success') loadFiles();
                } catch (error) {
                    console.error('Errore:', error);
                }
            }
            
            async function loadFiles() {
                try {
                    const response = await fetch('/api/documents');
                    const data = await response.json();
                    const filesList = document.getElementById('files-list');
                    filesList.innerHTML = data.documents.map(doc => `
                        <div class="file-item">
                            <span>📄 ${doc.filename}</span>
                            <button class="delete-btn" onclick="deleteFile('${doc.filename}')">✕</button>
                        </div>
                    `).join('');
                } catch (error) {
                    console.error('Errore:', error);
                }
            }
            
            async function deleteFile(filename) {
                try {
                    await fetch(`/api/documents/${filename}`, { method: 'DELETE' });
                    loadFiles();
                } catch (error) {
                    console.error('Errore:', error);
                }
            }
            
            loadFiles();
        </script>
    </body>
    </html>
    """
    return render_template_string(html)

@app.route('/api/status', methods=['GET'])
def status():
    return jsonify({
        "status": "Oracolo Covolo Online - Flask",
        "timestamp": datetime.now().isoformat()
    })

@app.route('/api/documents', methods=['GET'])
def list_documents():
    return jsonify({"documents": get_all_documents()})

@app.route('/api/upload', methods=['POST'])
def upload_document():
    if 'file' not in request.files:
        return jsonify({"status": "error", "message": "No file"}), 400
    
    file = request.files['file']
    if file.filename == '':
        return jsonify({"status": "error", "message": "No file"}), 400
    
    try:
        file_path = os.path.join(UPLOADS_DIR, file.filename)
        file.save(file_path)
        
        text_content = extract_text_from_file(file_path)
        file_size = os.path.getsize(file_path)
        saved = save_document_to_db(file.filename, text_content, file_size)
        
        if saved:
            return jsonify({"status": "success", "filename": file.filename})
        else:
            return jsonify({"status": "error", "message": "Save failed"}), 500
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500

@app.route('/api/ask', methods=['POST'])
def ask():
    data = request.get_json()
    question = (data.get('question') or "").strip()
    
    if not question:
        return jsonify({"error": "Empty question"}), 400
    
    try:
        matching_docs = search_documents(question)
        
        if not matching_docs:
            return jsonify({"answer": "Nessun documento trovato. Carica il listino Gessi!"})
        
        # Controlla macroregole
        parsed = {"product": question, "raw": question}
        incomp = macrorule_engine.check_incompatibility(parsed)
        if incomp:
            return jsonify({"answer": incomp.get("message", "Incompatibilità rilevata")})
        
        # Chiama DeepSeek
        doc_context = "\n".join([f"📄 {d['filename']}:\n{d['content'][:500]}" for d in matching_docs[:2]])
        
        prompt = f"""Tu sei consulente ESPERTO arredo bagno per Covolo.

DOCUMENTI:
{doc_context}

DOMANDA: {question}

Rispondi come narrativa calda, professionale e esperta."""

        response = httpx.post(
            DEEPSEEK_API_URL,
            headers={
                "Authorization": f"Bearer {DEEPSEEK_API_KEY}",
                "Content-Type": "application/json"
            },
            json={
                "model": DEEPSEEK_MODEL,
                "messages": [{
                    "role": "system",
                    "content": "Sei consulente professionista arredo bagno. Rispondi narrativa calda e esperta."
                }, {
                    "role": "user",
                    "content": prompt
                }],
                "temperature": 0.4,
                "max_tokens": 1500,
            },
            timeout=30
        )
        
        if response.status_code == 200:
            answer = response.json()["choices"][0]["message"]["content"]
            return jsonify({"answer": answer})
        else:
            return jsonify({"answer": "Errore nel generare risposta"})
    
    except Exception as e:
        return jsonify({"answer": f"Errore: {str(e)}"})

@app.route('/api/documents/<filename>', methods=['DELETE'])
def delete_document(filename):
    try:
        conn = sqlite3.connect(DB_PATH)
        c = conn.cursor()
        c.execute('DELETE FROM documents WHERE filename = ?', (filename,))
        conn.commit()
        conn.close()
        
        file_path = os.path.join(UPLOADS_DIR, filename)
        if os.path.exists(file_path):
            os.remove(file_path)
        
        return jsonify({"status": "success"})
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500

# ============================================================
# RUN
# ============================================================

if __name__ == '__main__':
    app.run(host='0.0.0.0', port=int(os.getenv('PORT', 5000)), debug=False)
